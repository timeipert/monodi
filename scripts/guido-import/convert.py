#!/usr/bin/env python3
"""
Converts a chant transcription typed in the GUIDO font (Dobszay/Bali, "guido6+", Hungarian
keyboard) inside a Word table into a Monodi-Zero import ZIP.

    python3 convert.py "R 1001.docx" [--out-dir DIR] [--virga flag|descending|normal]

Writes <stem>.monodi.zip (Sources > Import ZIP; one document per source, merged by id so a
re-import updates instead of duplicating) and <stem>.report.txt (what was read, what was not).

Layout of the .docx (one table per phrase of the chant)
  * a code row has the source siglum in column 0 and the GUIDO code in columns 1..N, one
    column per syllable; column 1 holds the clef;
  * the row right after a code row may be a text row (empty column 0, syllables in the
    columns). The first text row of a table is the shared text; a later one belongs to the code
    row above it and replaces the shared text in the columns it fills (a source's own spelling);
  * a run coloured BFBFBF is a copy of the reference row kept so the staff stays continuous:
    the source does NOT attest it. Only default-colour runs are read.

What a character means is read from glyph_table.json, which build_glyph_table.py measures out
of the font (every note glyph carries its own piece of staff, so its pitch is in the outline),
calibrated on the digits 0..9 = C4..E5 of the G clef. Meanings, with the guide section:
  digits            punctum                                    6.3
  q w e r t z u i o p  (+ õ ú î ...)  virga: stem down on the left   6.4, 2
  quilisma ladder (© « ¬ ...)  "vertically sliced oval" = two halves  2
  tenor ladder (¼ ½ ¾)  note framed by two strokes (tuba note)      2
  shift+digit hooks (' " + ! % / = ( ) §, Ö Ü Ó [ ])  plica       6.4
  upper-case row, a s d f ..., connecting strokes between notes     6.4, 2  (no pitch of their own)
  sharp/flat/natural signs by step                                 2 ("signs of alteration")
  , . : ;           phrasing marks (report only; --marks-in-text puts them into the lyrics)  2, 4.8
  -  _  Æ           empty staff: a gap of one notehead = a new neume   6.2, 4.6
  n                 empty glyph, 1/3 notehead wide: a short gap INSIDE a neume
  \\                empty glyph, 16/1000 em: nudges two notes together, keeps one ligature
Monodi representation (model.ts): syllable.notes = Spaced > NonSpaced (neume) > Grouped (ligature).
Pitch is absolute (step 0 = C4). virga -> note flag V (a stem down; --virga descending stores NoteType.Descending); quilisma ->
NoteType.Quilisma; plica -> note flag P, tenor note -> flag T, flat/sharp/natural -> NoteType
when the note is a plain punctum, else flag F/H/U (a note has one noteType). Flags are stored but
Monodi only draws flags that are defined under Settings > note flags: define V, P, T, F, H, U there.

Not decoded, and reported rather than guessed: the guide's liquescence tail, episema, prolongation
dot and reduced notehead are not identifiable in this font file (no glyph with that outline
exists beyond the hooks, which the guide itself calls plica), clefs other than G are only noted.
"""
import argparse
import json
import re
import sys
import uuid
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
GREY = 'BFBFBF'
JOIN, SHORT_GAP = '\\', 'n'
NOTE_CLASSES = ('punctum', 'virga', 'quilisma', 'tenor')
BASES = ['C', 'D', 'E', 'F', 'G', 'A', 'B']
NS = uuid.UUID('7b0f5d0e-6d0e-4f7e-9d57-3a1f4a0c6c11')

TABLE = json.loads(Path(__file__).with_name('glyph_table.json').read_text(encoding='utf-8'))['chars']


def step_to_pitch(step):
    di = 28 + step                      # C4 = 28 diatonic steps from C0
    return BASES[di % 7], di // 7


# ---------------------------------------------------------------------------------------
# decoding one syllable cell
# ---------------------------------------------------------------------------------------

@dataclass
class Note:
    step: int
    shape: str
    plica: bool = False
    accidental: str = ''


@dataclass
class Cell:
    neumes: list = field(default_factory=list)      # neume -> ligature group -> [Note]
    punct: str = ''
    clef: str = ''
    unhandled: list = field(default_factory=list)
    warnings: list = field(default_factory=list)

    def notes(self):
        return [n for ne in self.neumes for g in ne for n in g]


def decode_cell(code):
    cell = Cell()
    neumes, groups, notes = [], [], []
    pending = {}                         # step -> accidental waiting for its note
    last = None

    def close_group():
        nonlocal notes
        if notes:
            groups.append(notes)
            notes = []

    def close_neume():
        nonlocal groups
        close_group()
        if groups:
            neumes.append(groups)
            groups = []

    for ch in code:
        e = TABLE.get(ch)
        cls = e['cls'] if e else None
        if ch == JOIN or ch.isspace():
            continue
        if ch == SHORT_GAP:
            close_group()
        elif cls == 'gap':
            close_neume()
        elif cls in NOTE_CLASSES:
            note = Note(e['step'], cls, accidental=pending.pop(e['step'], ''))
            notes.append(note)
            last = note
        elif cls == 'plica':
            if last is not None and last.step == e['step']:
                last.plica = True
            else:
                cell.warnings.append(f'plica {ch!r} (step {e["step"]}) does not follow a note on its step')
        elif cls in ('sharp', 'flat', 'natural'):
            pending[e['step']] = cls
        elif cls == 'connector':
            continue
        elif cls == 'bar':
            cell.punct += ch
        elif cls == 'clef':
            cell.clef = e['clef']
        else:
            cell.unhandled.append(ch)
    close_neume()
    for step, acc in pending.items():
        cell.warnings.append(f'{acc} sign on step {step} found no note to apply to')
    cell.neumes = neumes
    return cell


# ---------------------------------------------------------------------------------------
# .docx extraction
# ---------------------------------------------------------------------------------------

def style_fonts(docx_path):
    with zipfile.ZipFile(docx_path) as z:
        xml = z.read('word/styles.xml').decode('utf-8')
    out = {}
    for m in re.finditer(r'<w:style [^>]*w:styleId="([^"]+)"[^>]*>(.*?)</w:style>', xml, re.S):
        fm = re.search(r'<w:rFonts[^>]*w:ascii="([^"]+)"', m.group(2))
        out[m.group(1)] = fm.group(1) if fm else None
    return out


def run_color(r):
    rpr = r.find(W + 'rPr')
    if rpr is not None:
        c = rpr.find(W + 'color')
        if c is not None:
            return (c.get(W + 'val') or '').upper()
    return None


def plain(cell):
    raw = ''.join(r.text for p in cell.paragraphs for r in p.runs if r.text)
    return re.sub(r'\s+', ' ', raw).strip()


def attested_code(cell):
    """(code of the default-colour runs, whether the cell has any run at all)."""
    parts, any_run = [], False
    for p in cell.paragraphs:
        for r in p.runs:
            if r.text:
                any_run = True
                if run_color(r._element) != GREY:
                    parts.append(r.text)
    return ''.join(parts), any_run


@dataclass
class Phrase:
    default_text: dict
    order: list
    codes: dict                  # siglum -> {col: code or None (all grey)}
    text_override: dict          # siglum -> {col: text}


def extract(docx_path):
    import docx
    d = docx.Document(docx_path)
    title = ''.join(r.text for r in d.paragraphs[0].runs) if d.paragraphs else ''
    phrases = []
    for t in d.tables:
        ph = Phrase({}, [], {}, {})
        ncols = len(t.rows[0].cells)
        first_text_done, last = False, None
        for row in t.rows:
            siglum = plain(row.cells[0])
            if siglum:
                if siglum not in ph.codes:
                    ph.order.append(siglum)
                    ph.codes[siglum], ph.text_override[siglum] = {}, {}
                for ci in range(1, ncols):
                    code, any_run = attested_code(row.cells[ci])
                    if any_run:
                        ph.codes[siglum][ci] = code if code else None
                last = siglum
            else:
                texts = {ci: plain(row.cells[ci]) for ci in range(1, ncols)}
                texts = {ci: x for ci, x in texts.items() if x}
                if not texts:
                    continue
                if not first_text_done:
                    ph.default_text.update(texts)
                    first_text_done = True
                elif last is not None:
                    ph.text_override[last].update(texts)
        phrases.append(ph)
    return title, phrases


# ---------------------------------------------------------------------------------------
# per-source assembly
# ---------------------------------------------------------------------------------------

@dataclass
class Syl:
    text: str
    cell: Cell


@dataclass
class SourceDoc:
    siglum: str
    syllables: list = field(default_factory=list)
    clef: str = ''
    warnings: list = field(default_factory=list)


def assemble(phrases, marks_in_text=False):
    order = []
    for ph in phrases:
        for s in ph.order:
            if s not in order:
                order.append(s)
    docs = {s: SourceDoc(s) for s in order}
    for pi, ph in enumerate(phrases, 1):
        for s in order:
            for ci in sorted(ph.codes.get(s, {})):
                code = ph.codes[s][ci]
                if code is None:
                    continue                                   # grey: not attested here
                text = ph.text_override.get(s, {}).get(ci) or ph.default_text.get(ci, '')
                cell = decode_cell(code)
                doc = docs[s]
                if cell.clef and not doc.clef:
                    doc.clef = cell.clef
                where = f'phrase {pi} col {ci} ({text or "no text"!r})'
                if cell.unhandled:
                    doc.warnings.append(f'{where}: not decoded {sorted(set(cell.unhandled))} in {code!r}')
                for w in cell.warnings:
                    doc.warnings.append(f'{where}: {w}')
                if not cell.neumes:
                    continue
                if not text:
                    doc.warnings.append(f'{where}: notes without text kept as an empty syllable')
                if marks_in_text and cell.punct and not text.endswith(cell.punct):
                    text += cell.punct
                doc.syllables.append(Syl(text, cell))
    return docs


# ---------------------------------------------------------------------------------------
# Monodi model (types/model.ts)
# ---------------------------------------------------------------------------------------

ACC_TYPE = {'flat': 'Flat', 'sharp': 'Sharp', 'natural': 'Natural'}
ACC_FLAG = {'flat': 'F', 'sharp': 'H', 'natural': 'U'}


class Ids:
    def __init__(self, seed):
        self.seed, self.n = seed, 0

    def next(self):
        self.n += 1
        return str(uuid.uuid5(NS, f'{self.seed}/{self.n}'))


def note_json(n, ids, virga):
    flags = []
    note_type = {'punctum': 'Normal', 'virga': 'Descending' if virga == 'descending' else 'Normal',
                 'quilisma': 'Quilisma', 'tenor': 'Normal'}[n.shape]
    if n.shape == 'virga' and virga == 'flag':
        flags.append('V')
    if n.shape == 'tenor':
        flags.append('T')
    if n.plica:
        flags.append('P')
    if n.accidental:
        if n.shape == 'punctum':
            note_type = ACC_TYPE[n.accidental]
        else:
            flags.append(ACC_FLAG[n.accidental])
    base, octave = step_to_pitch(n.step)
    out = {'uuid': ids.next(), 'noteType': note_type, 'base': base, 'liquescent': False,
           'octave': octave, 'focus': False}
    if flags:
        out['flags'] = flags
    return out


def syllable_json(s, ids, virga):
    return {
        'uuid': ids.next(), 'kind': 'Syllable', 'text': s.text, 'syllableType': 'Normal',
        'notes': {'spaced': [
            {'nonSpaced': [{'grouped': [note_json(n, ids, virga) for n in g]} for g in neume]}
            for neume in s.cell.neumes]},
    }


def root_json(doc, ids, virga):
    parts = []
    if doc.clef == 'G':
        parts.append({'kind': 'Clef', 'uuid': ids.next(), 'focus': False, 'base': 'G', 'octave': 4, 'shape': 'G'})
    parts += [syllable_json(s, ids, virga) for s in doc.syllables]
    return {
        'kind': 'RootContainer', 'uuid': ids.next(), 'comments': [], 'documentType': 'Level1',
        'children': [{'kind': 'MiscContainer', 'uuid': ids.next(),
                      'children': [{'kind': 'ZeileContainer', 'uuid': ids.next(), 'children': parts}]}],
    }


def parse_title(title):
    m = re.match(r'\s*(.*?)\s*_\s*(.*?)\s*$', title)
    return (m.group(1), m.group(2)) if m else ('', title.strip())


def safe(name):
    return re.sub(r'[\\/:*?"<>|]', '-', name).strip()


def write_zip(path, docs, title, virga, src_name):
    chant_id, incipit = parse_title(title)
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
        for siglum, doc in docs.items():
            if not doc.syllables:
                continue
            sid = safe(siglum)
            doc_id = safe(f'{siglum}-{chant_id}' if chant_id else f'{siglum}-{incipit}')
            z.writestr(f'{sid}/meta.json', json.dumps({'id': siglum, 'quellensigle': siglum}, ensure_ascii=False, indent=2))
            z.writestr(f'{sid}/{doc_id}/meta.json', json.dumps({
                'id': doc_id, 'quelle_id': siglum, 'textinitium': incipit,
                'custom': {k: v for k, v in (('chant_id', chant_id), ('import_file', src_name)) if v}},
                ensure_ascii=False, indent=2))
            z.writestr(f'{sid}/{doc_id}/data.json',
                       json.dumps(root_json(doc, Ids(doc_id), virga), ensure_ascii=False))


def describe(doc):
    def name(n):
        b, o = step_to_pitch(n.step)
        return f'{b}{o}{"~" if n.shape == "quilisma" else ""}{"^" if n.plica else ""}'
    lines = []
    for s in doc.syllables:
        neumes = ' | '.join(' / '.join(''.join(f'{name(n)} ' for n in g).strip() for g in ne) for ne in s.cell.neumes)
        marks = f'   [{s.cell.punct}]' if s.cell.punct else ''
        lines.append(f'  {s.text or "-":<12} {neumes}{marks}')
    return lines


def write_report(path, docs, title):
    with open(path, 'w', encoding='utf-8') as f:
        f.write(f'{title}\n\nnotation: " | " = neume boundary, " / " = short gap inside a neume, ~ quilisma, '
                f'^ plica, [x] phrasing mark after the syllable (Monodi has no place for it; see --marks-in-text)\n\n')
        for siglum, doc in docs.items():
            f.write(f'=== {siglum}: {len(doc.syllables)} syllables, clef {doc.clef or "none"}\n')
            f.writelines(line + '\n' for line in describe(doc))
            for w in doc.warnings:
                f.write(f'  ! {w}\n')
            f.write('\n')


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split('\n\n')[0])
    ap.add_argument('docx')
    ap.add_argument('--out-dir', default='.')
    ap.add_argument('--marks-in-text', action='store_true',
                    help='append the phrasing marks , . : ; to the syllable text (default: report only)')
    ap.add_argument('--virga', choices=['flag', 'descending', 'normal'], default='flag',
                    help='how the stemmed note is stored: flag V on a normal note (default), '
                         'NoteType.Descending, or not at all')
    a = ap.parse_args(argv)
    title, phrases = extract(a.docx)
    docs = assemble(phrases, a.marks_in_text)
    stem = Path(a.docx).stem
    out = Path(a.out_dir)
    out.mkdir(parents=True, exist_ok=True)
    write_zip(out / f'{stem}.monodi.zip', docs, title, a.virga, Path(a.docx).name)
    write_report(out / f'{stem}.report.txt', docs, title)
    kept = [s for s, d in docs.items() if d.syllables]
    print(f'{out / (stem + ".monodi.zip")}: {len(kept)} documents ({", ".join(kept)})')
    skipped = [s for s, d in docs.items() if not d.syllables]
    if skipped:
        print('nothing attested, skipped: ' + ', '.join(skipped))


if __name__ == '__main__':
    main()
